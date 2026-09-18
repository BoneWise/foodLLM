"""Read-only original workbook inventory for planning; no human/brain effects."""
import csv,hashlib,json,math,platform,subprocess,shutil
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import numpy as np
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/20260917_多营养OFC设计核对_v01'
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def rows(p):
 with Path(p).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def numeric(x):return isinstance(x,(int,float)) and math.isfinite(x)
def main():
 OUT.mkdir(exist_ok=True); assert not (OUT/'nutrition_inventory.json').exists()
 path=ROOT/'data/raw/metadata/food.pics.extended.database (Stand 2022).xlsx';before=sha(path)
 w=load_workbook(path,read_only=True,data_only=True);sheet=w['data'];r=list(sheet.iter_rows(values_only=True));headers=r[1]
 raw={str(int(v[0])):(i+3,v) for i,v in enumerate(r[2:]) if numeric(v[0])};assert len(raw)==sum(numeric(v[0]) for v in r[2:])
 prep_path=ROOT/'data/prepared/image_agreement_nutrition.csv';prepared=rows(prep_path); ids=[v['image_id'] for v in prepared]
 manifest_path=ROOT/'outputs/20260916_刺激与错配准备_v03/stimulus_manifest.csv';manifest=rows(manifest_path);primary=[v['image_id'] for v in manifest if v['P_eligible']=='True']
 perm_path=ROOT/'outputs/20260916_刺激与错配准备_v03/P_permutations.csv';perms=rows(perm_path)
 fields={'fibre_100g':8,'protein_100g':9,'fat_100g':10,'carbs_100g':11,'kcal_100g':12};report={};data_only={}
 for name,col in fields.items():
  values={i:raw[i][1][col] for i in ids};data_only[name]=values
  for rr in prepared:
   v=values[rr['image_id']]
   if numeric(v):assert abs(float(v)-float(rr[name]))<1e-9
   else: assert rr[name] in ['', 'nan','NaN']
  rec={'source_header':headers[col],'column':get_column_letter(col+1),'unit':'kcal/100 g' if name=='kcal_100g' else 'g/100 g','per100g_and_total_are_separate':True}
  for label,ss in [('full568',ids),('primary493',primary)]:
   vals=np.array([float(values[i]) for i in ss if numeric(values[i])]);missing=[i for i in ss if not numeric(values[i])]
   rec[label]={'available':len(vals),'missing_ids':missing,'zeros':int((vals==0).sum()),'range':[float(vals.min()),float(vals.max())],'q25_q75':[float(x) for x in np.quantile(vals,[.25,.75])]}
  precision=1 if name=='kcal_100g' else 2
  good=[p for p in perms if numeric(values[p['image_id']]) and numeric(values[p['donor_image_id']])]
  rec['reuse_kcal_P_diagnostic']={'eligible_pairs':len(good),'total_pairs':len(perms),'same_displayed_value_pairs':sum(f"{values[p['image_id']]:.{precision}f}"==f"{values[p['donor_image_id']]:.{precision}f}" for p in good),'precision':precision,'new_permutations_generated':False}
  report[name]=rec
 complete=[i for i in primary if all(numeric(data_only[f][i]) for f in fields)]
 arr=np.array([[data_only[f][i] for f in fields] for i in complete],dtype=float)
 ranked=[]
 for column in arr.T:
  _,inverse,count=np.unique(column,return_inverse=True,return_counts=True)
  ranked.append((np.cumsum(count)-(count-1)/2)[inverse])
 corr=np.corrcoef(ranked)
 out={'status':'read_only_inventory_complete','source':str(path.relative_to(ROOT)),'sheet':'data','header_row':2,'data_rows':[3,1215],'dimensions':report,'current493_complete_all_five':len(complete),'full568':len(ids),'no_independent_sugar_sodium_or_vitamin_field':True,'other_columns_are_classification_image_characteristics_or_normative_ratings':True,'existing_prepared_values_match_raw':True,'original_workbook_unchanged':before==sha(path),'nutrient_only_Spearman':{'sample':'current493_complete_cases','n':len(complete),'order':list(fields),'matrix':corr.tolist()},'source_hashes':{str(p.relative_to(ROOT)):sha(p) for p in [path,prep_path,manifest_path,perm_path,Path(__file__)]},'versions':{'python':platform.python_version(),'numpy':np.__version__},'git_base':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),'model_or_neural_effects_computed':False,'new_P_or_ROI_masks_created':False,'seed':None}
 (OUT/'nutrition_inventory.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n');shutil.copy2(__file__,OUT/'script_snapshot.py')
 print(json.dumps({'available493':{f:report[f]['primary493']['available'] for f in fields},'unchanged':out['original_workbook_unchanged'],'existingP_equal':{f:report[f]['reuse_kcal_P_diagnostic']['same_displayed_value_pairs'] for f in fields}},ensure_ascii=False))
if __name__=='__main__':main()
